from __future__ import annotations

from pathlib import Path
from typing import Iterable

from openpyxl import load_workbook

from backend.app.services.reports.excel_export import SHEET_NAMES


def validate_excel_workbook(
    workbook_path: str | Path,
    key_sheets_with_data: Iterable[str] = (),
) -> list[str]:
    path = Path(workbook_path)
    issues: list[str] = []

    if not path.exists():
        return [f"Workbook does not exist: {path}"]

    workbook = load_workbook(path, read_only=False)

    missing_sheets = [sheet for sheet in SHEET_NAMES if sheet not in workbook.sheetnames]
    if missing_sheets:
        issues.append("Missing expected sheet(s): " + ", ".join(missing_sheets))

    if "Executive Summary" in workbook.sheetnames:
        summary = workbook["Executive Summary"]
        if summary.max_row < 2 or not summary["A1"].value:
            issues.append("Executive Summary is empty.")

    if "Assumptions" in workbook.sheetnames:
        assumptions = workbook["Assumptions"]
        if assumptions.max_row < 2 or not assumptions["A2"].value:
            issues.append("Assumptions sheet is empty.")

    for sheet_name in key_sheets_with_data:
        if sheet_name not in workbook.sheetnames:
            issues.append(f"Cannot validate missing sheet: {sheet_name}")
            continue
        sheet = workbook[sheet_name]
        if sheet.max_row < 2 or sheet["A2"].value in (None, ""):
            issues.append(f"Expected at least one data row in {sheet_name}.")

    return issues

