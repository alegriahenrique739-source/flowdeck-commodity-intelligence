from __future__ import annotations

from copy import copy
from datetime import UTC, datetime
from decimal import Decimal
from math import ceil, floor
from pathlib import Path
from typing import Any, Iterable

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from backend.app.services.curves import ForwardCurve
from backend.app.services.positions import (
    signed_exposure_bbl,
    signed_physical_exposure_bbl,
)
from backend.app.services.reports.excel_schemas import (
    CalendarSpreadReportRow,
    ExcelReportInput,
)


REPORT_TITLE = "FlowDeck Commodity Trading Intelligence Report"
SHEET_NAMES = (
    "Executive Summary",
    "Market Curves",
    "Calendar Spreads",
    "Futures Positions",
    "Physical Cargoes",
    "Net Exposure",
    "Hedge Simulation",
    "Stress P&L",
    "Assumptions",
)
DEFAULT_ASSUMPTIONS = (
    "Synthetic demo data only.",
    "Analytics only.",
    "Not trade execution.",
    "Hedge recommendations require human review.",
    "No licensed market data redistribution.",
    "BUY physical = long.",
    "SELL physical = short.",
    "BUY futures = long.",
    "SELL futures = short.",
    "Positive net exposure = long price exposure.",
    "Negative net exposure = short price exposure.",
)

HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
TITLE_FILL = PatternFill("solid", fgColor="D9EAF7")
WHITE_FONT = Font(color="FFFFFF", bold=True)
BOLD_FONT = Font(bold=True)
TITLE_FONT = Font(bold=True, size=14, color="1F4E78")
NOTE_FONT = Font(italic=True, color="666666")
GREEN_FILL = PatternFill("solid", fgColor="E2F0D9")
RED_FILL = PatternFill("solid", fgColor="FCE4D6")
SECTION_FILL = PatternFill("solid", fgColor="EAF2F8")
EXPOSURE_FILL = PatternFill("solid", fgColor="FFF2CC")

BARRELS_FORMAT = "#,##0"
LOTS_FORMAT = "#,##0.00"
USD_PER_BBL_FORMAT = "$#,##0.00"
PNL_FORMAT = "$#,##0;[Red]-$#,##0"


def build_excel_report(
    report_input: ExcelReportInput, output_path: str | Path
) -> Path:
    path = Path(output_path)
    path.parent.mkdir(parents=True, exist_ok=True)

    workbook = Workbook()
    workbook.remove(workbook.active)

    summary_sheet = workbook.create_sheet("Executive Summary")
    market_sheet = workbook.create_sheet("Market Curves")
    spreads_sheet = workbook.create_sheet("Calendar Spreads")
    futures_sheet = workbook.create_sheet("Futures Positions")
    physical_sheet = workbook.create_sheet("Physical Cargoes")
    net_sheet = workbook.create_sheet("Net Exposure")
    hedge_sheet = workbook.create_sheet("Hedge Simulation")
    stress_sheet = workbook.create_sheet("Stress P&L")
    assumptions_sheet = workbook.create_sheet("Assumptions")

    _write_executive_summary(summary_sheet, report_input)
    _write_table_sheet(
        market_sheet,
        _market_data_headers(),
        _market_data_rows(report_input),
        empty_note="No market data rows supplied.",
    )
    _write_table_sheet(
        spreads_sheet,
        _calendar_spread_headers(),
        _calendar_spread_rows(report_input),
        empty_note="No calendar spreads supplied.",
    )
    _write_table_sheet(
        futures_sheet,
        _futures_headers(),
        _futures_rows(report_input),
        empty_note="No futures positions supplied.",
    )
    _write_table_sheet(
        physical_sheet,
        _physical_headers(),
        _physical_rows(report_input),
        empty_note="No physical cargoes supplied.",
    )
    _write_table_sheet(
        net_sheet,
        _net_exposure_headers(),
        _net_exposure_rows(report_input),
        empty_note="No net exposure result supplied.",
        exposure_columns={"net_exposure_bbl", "absolute_net_exposure_bbl"},
    )
    _write_table_sheet(
        hedge_sheet,
        _hedge_headers(),
        _hedge_rows(report_input),
        empty_note="No hedge simulation result supplied.",
    )
    _write_table_sheet(
        stress_sheet,
        _stress_headers(),
        _stress_rows(report_input),
        empty_note="No stress P&L result supplied.",
        pnl_columns={"stress_pnl_usd"},
    )
    _write_assumptions(assumptions_sheet, report_input.assumptions)

    _add_market_curve_chart(market_sheet)
    _add_net_exposure_chart(net_sheet)
    _add_stress_pnl_chart(stress_sheet)

    for sheet in workbook.worksheets:
        _apply_common_sheet_formatting(sheet)

    workbook.save(path)
    return path


def _write_executive_summary(sheet: Worksheet, report_input: ExcelReportInput) -> None:
    sheet["A1"] = REPORT_TITLE
    sheet["A1"].font = TITLE_FONT
    sheet["A1"].fill = TITLE_FILL
    sheet.merge_cells("A1:E1")

    sheet["A3"] = "Generated timestamp"
    sheet["B3"] = datetime.now(UTC).replace(microsecond=0).isoformat()
    sheet["A4"] = "Scope"
    sheet["B4"] = "Oil / Brent-WTI / synthetic demo data"
    sheet["A5"] = "Note"
    sheet["B5"] = "Analytics only. Hedge recommendations require human review."

    metrics = _summary_metrics(report_input)
    _write_section_header(sheet, 7, "Key Metrics")
    sheet["A8"] = "Metric"
    sheet["B8"] = "Value"
    for cell in sheet[8]:
        cell.fill = HEADER_FILL
        cell.font = WHITE_FONT

    for row_number, (metric, value) in enumerate(metrics, start=9):
        sheet.cell(row=row_number, column=1, value=metric)
        sheet.cell(row=row_number, column=2, value=_excel_value(value))

    risk_start = 9 + len(metrics) + 2
    _write_section_header(sheet, risk_start, "Risk and Hedge Summary")
    risk_rows = _risk_hedge_summary(report_input)
    sheet.cell(row=risk_start + 1, column=1, value="Metric")
    sheet.cell(row=risk_start + 1, column=2, value="Value")
    for cell in sheet[risk_start + 1]:
        cell.fill = HEADER_FILL
        cell.font = WHITE_FONT
    for row_number, (metric, value) in enumerate(risk_rows, start=risk_start + 2):
        sheet.cell(row=row_number, column=1, value=metric)
        sheet.cell(row=row_number, column=2, value=_excel_value(value))

    assumptions_start = risk_start + len(risk_rows) + 4
    _write_section_header(sheet, assumptions_start, "Assumptions Note")
    sheet.cell(
        row=assumptions_start + 1,
        column=1,
        value=(
            "Synthetic demo data only; no licensed market data redistribution; "
            "not trade execution or financial advice."
        ),
    )
    sheet.merge_cells(
        start_row=assumptions_start + 1,
        start_column=1,
        end_row=assumptions_start + 1,
        end_column=5,
    )
    sheet.freeze_panes = "A8"
    _format_numeric_columns(sheet, {"Value": "general"})


def _write_section_header(sheet: Worksheet, row_number: int, label: str) -> None:
    sheet.cell(row=row_number, column=1, value=label)
    sheet.cell(row=row_number, column=1).font = BOLD_FONT
    sheet.cell(row=row_number, column=1).fill = SECTION_FILL
    sheet.merge_cells(
        start_row=row_number,
        start_column=1,
        end_row=row_number,
        end_column=5,
    )


def _summary_metrics(report_input: ExcelReportInput) -> list[tuple[str, Any]]:
    net_summary = (
        report_input.net_exposure_result.summary
        if report_input.net_exposure_result is not None
        else None
    )
    hedge_summary = (
        report_input.hedge_simulation_result.summary
        if report_input.hedge_simulation_result is not None
        else None
    )
    stress_summary = report_input.stress_result.summary if report_input.stress_result else None

    return [
        ("Number of market data rows", len(report_input.market_data_rows)),
        ("Number of curve summaries", len(report_input.curve_summaries)),
        ("Number of futures positions", len(report_input.futures_positions)),
        ("Number of physical cargoes", len(report_input.physical_cargoes)),
        (
            "Number of net exposure buckets",
            net_summary.number_of_exposure_buckets if net_summary else 0,
        ),
        ("Total net exposure bbl", net_summary.total_net_exposure_bbl if net_summary else 0),
        (
            "Total absolute net exposure bbl",
            net_summary.total_absolute_net_exposure_bbl if net_summary else 0,
        ),
        (
            "Total gross component exposure bbl",
            net_summary.total_gross_component_exposure_bbl if net_summary else 0,
        ),
        (
            "Total post-hedge absolute net exposure bbl",
            hedge_summary.total_post_hedge_absolute_net_exposure_bbl if hedge_summary else 0,
        ),
        (
            "Total stress P&L USD",
            stress_summary.total_stress_pnl_usd if stress_summary else 0,
        ),
    ]


def _risk_hedge_summary(report_input: ExcelReportInput) -> list[tuple[str, Any]]:
    net_summary = (
        report_input.net_exposure_result.summary
        if report_input.net_exposure_result is not None
        else None
    )
    hedge_summary = (
        report_input.hedge_simulation_result.summary
        if report_input.hedge_simulation_result is not None
        else None
    )
    stress_summary = report_input.stress_result.summary if report_input.stress_result else None

    return [
        (
            "Largest long bucket",
            net_summary.largest_long_bucket.label
            if net_summary and net_summary.largest_long_bucket
            else "n/a",
        ),
        (
            "Largest short bucket",
            net_summary.largest_short_bucket.label
            if net_summary and net_summary.largest_short_bucket
            else "n/a",
        ),
        (
            "Hedge recommendations",
            hedge_summary.number_of_recommendations if hedge_summary else 0,
        ),
        (
            "No-action hedge buckets",
            hedge_summary.number_of_no_action_buckets if hedge_summary else 0,
        ),
        (
            "Stress scenario",
            stress_summary.scenario_name if stress_summary else "n/a",
        ),
        (
            "Total stress gain USD",
            stress_summary.total_gain_usd if stress_summary else 0,
        ),
        (
            "Total stress loss USD",
            stress_summary.total_loss_usd if stress_summary else 0,
        ),
    ]


def _write_table_sheet(
    sheet: Worksheet,
    headers: list[str],
    rows: list[list[Any]],
    empty_note: str,
    pnl_columns: set[str] | None = None,
    exposure_columns: set[str] | None = None,
) -> None:
    for column_number, header in enumerate(headers, start=1):
        cell = sheet.cell(row=1, column=column_number, value=header)
        cell.fill = HEADER_FILL
        cell.font = WHITE_FONT
        cell.alignment = Alignment(horizontal="center")

    if rows:
        for row_number, row in enumerate(rows, start=2):
            for column_number, value in enumerate(row, start=1):
                sheet.cell(row=row_number, column=column_number, value=_excel_value(value))
        sheet.auto_filter.ref = sheet.dimensions
    else:
        sheet.cell(row=2, column=1, value=empty_note)
        sheet.cell(row=2, column=1).font = NOTE_FONT

    sheet.freeze_panes = "A2"
    _format_numeric_columns(sheet, _number_format_map(headers))
    if pnl_columns:
        _apply_pnl_conditional_formatting(sheet, headers, pnl_columns)
    if exposure_columns:
        _apply_exposure_conditional_formatting(sheet, headers, exposure_columns)


def _write_assumptions(sheet: Worksheet, assumptions: Iterable[str]) -> None:
    sheet["A1"] = "Assumptions and Controls"
    sheet["A1"].fill = HEADER_FILL
    sheet["A1"].font = WHITE_FONT
    notes = tuple(assumptions) if assumptions else DEFAULT_ASSUMPTIONS
    for row_number, note in enumerate(notes, start=2):
        sheet.cell(row=row_number, column=1, value=note)
    sheet.freeze_panes = "A2"


def _market_data_headers() -> list[str]:
    return [
        "valuation_date",
        "commodity",
        "contract_code",
        "contract_month",
        "expiry_date",
        "price",
        "currency",
        "unit",
        "source",
    ]


def _market_data_rows(report_input: ExcelReportInput) -> list[list[Any]]:
    return [
        [
            row.valuation_date,
            row.commodity,
            row.contract_code,
            row.contract_month,
            row.expiry_date,
            row.price,
            row.currency,
            row.unit,
            row.source,
        ]
        for row in report_input.market_data_rows
    ]


def _calendar_spread_headers() -> list[str]:
    return [
        "valuation_date",
        "commodity",
        "spread_name",
        "near_contract",
        "far_contract",
        "spread_usd_per_bbl",
    ]


def _calendar_spread_rows(report_input: ExcelReportInput) -> list[list[Any]]:
    spreads = report_input.calendar_spreads or _spreads_from_curves(
        report_input.curve_summaries
    )
    return [
        [
            spread.valuation_date,
            spread.commodity,
            spread.spread_name,
            spread.near_contract,
            spread.far_contract,
            spread.spread_usd_per_bbl,
        ]
        for spread in spreads
    ]


def _spreads_from_curves(
    curves: tuple[ForwardCurve, ...]
) -> tuple[CalendarSpreadReportRow, ...]:
    rows: list[CalendarSpreadReportRow] = []
    for curve in curves:
        for name, spread in curve.calendar_spreads.items():
            rows.append(
                CalendarSpreadReportRow(
                    valuation_date=curve.valuation_date,
                    commodity=curve.commodity,
                    spread_name=name,
                    near_contract=spread.near_contract,
                    far_contract=spread.far_contract,
                    spread_usd_per_bbl=spread.spread,
                )
            )
    return tuple(rows)


def _futures_headers() -> list[str]:
    return [
        "position_id",
        "trade_date",
        "book",
        "commodity",
        "contract_code",
        "contract_month",
        "direction",
        "lots",
        "contract_size",
        "entry_price",
        "signed_exposure_bbl",
    ]


def _futures_rows(report_input: ExcelReportInput) -> list[list[Any]]:
    return [
        [
            row.position_id,
            row.trade_date,
            row.book,
            row.commodity,
            row.contract_code,
            row.contract_month,
            row.direction,
            row.lots,
            row.contract_size,
            row.entry_price,
            signed_exposure_bbl(row),
        ]
        for row in report_input.futures_positions
    ]


def _physical_headers() -> list[str]:
    return [
        "cargo_id",
        "trade_date",
        "book",
        "commodity",
        "buy_sell",
        "volume_bbl",
        "pricing_index",
        "pricing_month",
        "delivery_start",
        "delivery_end",
        "location",
        "fixed_price",
        "differential",
        "freight_cost_per_bbl",
        "storage_cost_per_bbl",
        "signed_physical_exposure_bbl",
    ]


def _physical_rows(report_input: ExcelReportInput) -> list[list[Any]]:
    return [
        [
            row.cargo_id,
            row.trade_date,
            row.book,
            row.commodity,
            row.buy_sell,
            row.volume_bbl,
            row.pricing_index,
            row.pricing_month,
            row.delivery_start,
            row.delivery_end,
            row.location,
            row.fixed_price,
            row.differential,
            row.freight_cost_per_bbl,
            row.storage_cost_per_bbl,
            signed_physical_exposure_bbl(row),
        ]
        for row in report_input.physical_cargoes
    ]


def _net_exposure_headers() -> list[str]:
    return [
        "hedge_index",
        "exposure_month",
        "book",
        "physical_exposure_bbl",
        "paper_exposure_bbl",
        "net_exposure_bbl",
        "absolute_net_exposure_bbl",
        "gross_component_exposure_bbl",
        "label",
    ]


def _net_exposure_rows(report_input: ExcelReportInput) -> list[list[Any]]:
    if report_input.net_exposure_result is None:
        return []
    return [
        [
            bucket.hedge_index,
            bucket.exposure_month,
            bucket.book,
            bucket.physical_exposure_bbl,
            bucket.paper_exposure_bbl,
            bucket.net_exposure_bbl,
            bucket.absolute_net_exposure_bbl,
            bucket.gross_component_exposure_bbl,
            bucket.label,
        ]
        for bucket in report_input.net_exposure_result.buckets
    ]


def _hedge_headers() -> list[str]:
    return [
        "hedge_index",
        "exposure_month",
        "book",
        "current_net_exposure_bbl",
        "target_residual_exposure_bbl",
        "required_hedge_exposure_bbl",
        "recommended_direction",
        "recommended_lots_exact",
        "recommended_lots_rounded",
        "rounded_hedge_exposure_bbl",
        "post_hedge_net_exposure_bbl",
        "residual_difference_vs_target_bbl",
        "recommendation_label",
    ]


def _hedge_rows(report_input: ExcelReportInput) -> list[list[Any]]:
    if report_input.hedge_simulation_result is None:
        return []
    return [
        [
            row.hedge_index,
            row.exposure_month,
            row.book,
            row.current_net_exposure_bbl,
            row.target_residual_exposure_bbl,
            row.required_hedge_exposure_bbl,
            row.recommended_direction,
            row.recommended_lots_exact,
            row.recommended_lots_rounded,
            row.rounded_hedge_exposure_bbl,
            row.post_hedge_net_exposure_bbl,
            row.residual_difference_vs_target_bbl,
            row.recommendation_label,
        ]
        for row in report_input.hedge_simulation_result.recommendations
    ]


def _stress_headers() -> list[str]:
    return [
        "scenario_name",
        "hedge_index",
        "exposure_month",
        "book",
        "net_exposure_bbl",
        "price_delta_usd_per_bbl",
        "stress_pnl_usd",
        "pnl_direction",
        "label",
    ]


def _stress_rows(report_input: ExcelReportInput) -> list[list[Any]]:
    if report_input.stress_result is None:
        return []
    return [
        [
            row.scenario_name,
            row.hedge_index,
            row.exposure_month,
            row.book,
            row.net_exposure_bbl,
            row.price_delta_usd_per_bbl,
            row.stress_pnl_usd,
            row.pnl_direction,
            row.label,
        ]
        for row in report_input.stress_result.bucket_results
    ]


def _excel_value(value: Any) -> Any:
    if isinstance(value, Decimal):
        return float(value)
    if hasattr(value, "value"):
        return value.value
    return value


def _number_format_map(headers: list[str]) -> dict[str, str]:
    formats: dict[str, str] = {}
    for header in headers:
        if header.endswith("_bbl") or header == "contract_size":
            formats[header] = BARRELS_FORMAT
        if "lots" in header:
            formats[header] = LOTS_FORMAT
        if header in {
            "price",
            "entry_price",
            "fixed_price",
            "differential",
            "freight_cost_per_bbl",
            "storage_cost_per_bbl",
            "spread_usd_per_bbl",
            "price_delta_usd_per_bbl",
        }:
            formats[header] = USD_PER_BBL_FORMAT
        if header.endswith("_usd") or header == "stress_pnl_usd":
            formats[header] = PNL_FORMAT
    return formats


def _format_numeric_columns(sheet: Worksheet, formats: dict[str, str]) -> None:
    headers = [cell.value for cell in sheet[1]]
    for column_number, header in enumerate(headers, start=1):
        number_format = formats.get(str(header))
        if not number_format:
            continue
        for cell in sheet.iter_cols(
            min_col=column_number,
            max_col=column_number,
            min_row=2,
            max_row=sheet.max_row,
        ):
            for item in cell:
                item.number_format = number_format


def _apply_pnl_conditional_formatting(
    sheet: Worksheet, headers: list[str], pnl_columns: set[str]
) -> None:
    for index, header in enumerate(headers, start=1):
        if header not in pnl_columns or sheet.max_row < 2:
            continue
        column = get_column_letter(index)
        cell_range = f"{column}2:{column}{sheet.max_row}"
        sheet.conditional_formatting.add(
            cell_range,
            CellIsRule(operator="lessThan", formula=["0"], fill=RED_FILL),
        )
        sheet.conditional_formatting.add(
            cell_range,
            CellIsRule(operator="greaterThan", formula=["0"], fill=GREEN_FILL),
        )


def _apply_exposure_conditional_formatting(
    sheet: Worksheet, headers: list[str], exposure_columns: set[str]
) -> None:
    for index, header in enumerate(headers, start=1):
        if header not in exposure_columns or sheet.max_row < 2:
            continue
        column = get_column_letter(index)
        cell_range = f"{column}2:{column}{sheet.max_row}"
        sheet.conditional_formatting.add(
            cell_range,
            CellIsRule(operator="greaterThan", formula=["0"], fill=EXPOSURE_FILL),
        )
        sheet.conditional_formatting.add(
            cell_range,
            CellIsRule(operator="lessThan", formula=["0"], fill=EXPOSURE_FILL),
        )


def _add_market_curve_chart(sheet: Worksheet) -> None:
    if sheet.max_row < 3:
        return

    helper_start_col = sheet.max_column + 2
    rows = list(sheet.iter_rows(min_row=2, values_only=True))
    months = sorted({row[3] for row in rows if row[3]})
    commodities = sorted({row[1] for row in rows if row[1]})
    if not months or not commodities:
        return

    price_by_key = {(row[1], row[3]): row[5] for row in rows}
    prices = [
        float(row[5])
        for row in rows
        if row[5] is not None and isinstance(row[5], (int, float))
    ]
    sheet.cell(row=1, column=helper_start_col, value="contract_month")
    for offset, commodity in enumerate(commodities, start=1):
        sheet.cell(row=1, column=helper_start_col + offset, value=commodity)
    for row_number, month in enumerate(months, start=2):
        sheet.cell(row=row_number, column=helper_start_col, value=month)
        for offset, commodity in enumerate(commodities, start=1):
            sheet.cell(
                row=row_number,
                column=helper_start_col + offset,
                value=price_by_key.get((commodity, month)),
            )

    chart = LineChart()
    chart.title = "Forward Curves"
    chart.y_axis.title = "USD/bbl"
    chart.x_axis.title = "Contract Month"
    chart.style = 13
    data = Reference(
        sheet,
        min_col=helper_start_col + 1,
        max_col=helper_start_col + len(commodities),
        min_row=1,
        max_row=len(months) + 1,
    )
    categories = Reference(
        sheet,
        min_col=helper_start_col,
        min_row=2,
        max_row=len(months) + 1,
    )
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(categories)
    if prices:
        chart.y_axis.scaling.min = floor(min(prices) - 1)
        chart.y_axis.scaling.max = ceil(max(prices) + 1)
    for series in chart.series:
        series.marker.symbol = "circle"
        series.marker.size = 5
    chart.height = 7
    chart.width = 15
    sheet.add_chart(chart, f"{get_column_letter(helper_start_col)}8")

    for col in range(helper_start_col, helper_start_col + len(commodities) + 1):
        sheet.column_dimensions[get_column_letter(col)].hidden = True


def _add_net_exposure_chart(sheet: Worksheet) -> None:
    _add_bar_chart(
        sheet=sheet,
        title="Net Exposure by Bucket",
        value_header="net_exposure_bbl",
        category_header="label",
        y_axis_title="bbl",
        anchor="K2",
    )


def _add_stress_pnl_chart(sheet: Worksheet) -> None:
    _add_bar_chart(
        sheet=sheet,
        title="Stress P&L by Bucket",
        value_header="stress_pnl_usd",
        category_header="label",
        y_axis_title="USD",
        anchor="K2",
    )


def _add_bar_chart(
    sheet: Worksheet,
    title: str,
    value_header: str,
    category_header: str,
    y_axis_title: str,
    anchor: str,
) -> None:
    if sheet.max_row < 2:
        return
    headers = [cell.value for cell in sheet[1]]
    if value_header not in headers or category_header not in headers:
        return

    value_col = headers.index(value_header) + 1
    category_col = headers.index(category_header) + 1
    chart = BarChart()
    chart.type = "bar"
    chart.grouping = "clustered"
    chart.overlap = 0
    chart.title = title
    chart.x_axis.title = y_axis_title
    chart.y_axis.title = "Bucket"
    chart.style = 10
    data = Reference(sheet, min_col=value_col, min_row=1, max_row=sheet.max_row)
    categories = Reference(sheet, min_col=category_col, min_row=2, max_row=sheet.max_row)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(categories)
    chart.legend = None
    chart.height = 8
    chart.width = 17
    sheet.add_chart(chart, anchor)


def _apply_common_sheet_formatting(sheet: Worksheet) -> None:
    sheet.sheet_view.showGridLines = False
    for row in sheet.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=False)

    for column_cells in sheet.columns:
        max_length = 0
        column_letter = get_column_letter(column_cells[0].column)
        if sheet.column_dimensions[column_letter].hidden:
            continue
        for cell in column_cells:
            max_length = max(max_length, len(str(cell.value or "")))
        sheet.column_dimensions[column_letter].width = min(max(max_length + 2, 12), 42)

    if sheet.max_row > 1:
        for cell in sheet[1]:
            updated_font = copy(cell.font)
            updated_font.bold = True
            cell.font = updated_font
