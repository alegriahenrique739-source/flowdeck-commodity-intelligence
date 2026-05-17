from pathlib import Path
from uuid import uuid4

from openpyxl import load_workbook

from backend.app.services.reports import (
    ExcelReportInput,
    build_excel_report,
    validate_excel_workbook,
)
from scripts.build_sample_excel_report import build_sample_report


EXPECTED_SHEETS = [
    "Executive Summary",
    "Market Curves",
    "Calendar Spreads",
    "Futures Positions",
    "Physical Cargoes",
    "Net Exposure",
    "Hedge Simulation",
    "Stress P&L",
    "Assumptions",
]


def test_excel_report_file_is_created() -> None:
    output_path = build_excel_report(ExcelReportInput(), _test_output_path())

    assert output_path.exists()
    assert output_path.stat().st_size > 0


def test_workbook_contains_all_expected_sheet_names() -> None:
    output_path = build_excel_report(ExcelReportInput(), _test_output_path())

    workbook = load_workbook(output_path, read_only=False)

    assert workbook.sheetnames == EXPECTED_SHEETS


def test_workbook_is_not_empty() -> None:
    output_path = build_excel_report(ExcelReportInput(), _test_output_path())

    workbook = load_workbook(output_path, read_only=False)
    summary = workbook["Executive Summary"]

    assert summary.max_row >= 10
    assert summary["A1"].value == "FlowDeck Commodity Trading Intelligence Report"


def test_executive_summary_and_assumptions_sheets_exist() -> None:
    output_path = build_excel_report(ExcelReportInput(), _test_output_path())

    workbook = load_workbook(output_path, read_only=True)

    assert "Executive Summary" in workbook.sheetnames
    assert "Assumptions" in workbook.sheetnames
    assert workbook["Assumptions"]["A2"].value == "Synthetic demo data only."


def test_missing_optional_sections_do_not_crash() -> None:
    output_path = build_excel_report(ExcelReportInput(), _test_output_path())

    workbook = load_workbook(output_path, read_only=True)

    assert workbook["Market Curves"]["A2"].value == "No market data rows supplied."
    assert workbook["Net Exposure"]["A2"].value == "No net exposure result supplied."
    assert workbook["Stress P&L"]["A2"].value == "No stress P&L result supplied."


def test_sample_report_builder_creates_valid_workbook() -> None:
    output_path = build_sample_report(_test_output_path("flowdeck_sample_report"))

    workbook = load_workbook(output_path, read_only=True)

    assert workbook.sheetnames == EXPECTED_SHEETS
    assert workbook["Market Curves"].max_row > 1
    assert workbook["Futures Positions"].max_row > 1
    assert workbook["Physical Cargoes"].max_row > 1
    assert workbook["Net Exposure"].max_row > 1
    assert workbook["Hedge Simulation"].max_row > 1
    assert workbook["Stress P&L"].max_row > 1


def test_charts_are_present_on_relevant_sheets() -> None:
    output_path = build_sample_report(_test_output_path("flowdeck_sample_report"))

    workbook = load_workbook(output_path, read_only=False)

    assert len(workbook["Market Curves"]._charts) >= 1
    assert len(workbook["Net Exposure"]._charts) >= 1
    assert len(workbook["Stress P&L"]._charts) >= 1


def test_chart_series_logic_is_professional() -> None:
    output_path = build_sample_report(_test_output_path("flowdeck_sample_report"))

    workbook = load_workbook(output_path, read_only=False)
    market_chart = workbook["Market Curves"]._charts[0]
    net_chart = workbook["Net Exposure"]._charts[0]
    stress_chart = workbook["Stress P&L"]._charts[0]

    assert len(market_chart.series) == 2
    assert market_chart.y_axis.scaling.min is not None
    assert market_chart.y_axis.scaling.max is not None
    assert len(net_chart.series) == 1
    assert net_chart.type == "bar"
    assert net_chart.legend is None
    assert len(stress_chart.series) == 1
    assert stress_chart.type == "bar"
    assert stress_chart.legend is None


def test_freeze_panes_are_set_on_tabular_sheets() -> None:
    output_path = build_sample_report(_test_output_path("flowdeck_sample_report"))

    workbook = load_workbook(output_path, read_only=False)

    for sheet_name in EXPECTED_SHEETS:
        if sheet_name == "Executive Summary":
            continue
        assert workbook[sheet_name].freeze_panes == "A2"


def test_report_qa_helper_validates_sample_workbook() -> None:
    output_path = build_sample_report(_test_output_path("flowdeck_sample_report"))

    issues = validate_excel_workbook(
        output_path,
        key_sheets_with_data=[
            "Market Curves",
            "Futures Positions",
            "Physical Cargoes",
            "Net Exposure",
            "Hedge Simulation",
            "Stress P&L",
        ],
    )

    assert issues == []


def _test_output_path(prefix: str = "report") -> Path:
    output_dir = Path(__file__).resolve().parents[2] / "reports" / "test_outputs"
    return output_dir / f"{prefix}_{uuid4().hex}.xlsx"
