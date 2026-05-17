from pathlib import Path

from openpyxl import load_workbook

from scripts.run_demo_workflow import run_demo_workflow


PROJECT_ROOT = Path(__file__).resolve().parents[2]


def test_api_contract_doc_exists() -> None:
    path = PROJECT_ROOT / "docs" / "API_CONTRACT.md"

    assert path.exists()
    assert "GET /health" in path.read_text()


def test_demo_guide_exists() -> None:
    path = PROJECT_ROOT / "docs" / "DEMO_GUIDE.md"

    assert path.exists()
    assert "python -m uvicorn backend.app.main:app --reload" in path.read_text()


def test_run_demo_workflow_creates_excel_file() -> None:
    output_root = PROJECT_ROOT / "reports" / "test_outputs" / "demo_workflow"
    output_root.mkdir(parents=True, exist_ok=True)
    output_path = output_root / "demo_workflow_report.xlsx"
    runs_root = output_root / "runs"

    summary = run_demo_workflow(output_path, runs_root=runs_root)

    assert output_path.exists()
    assert Path(summary["metadata_output_path"]).exists()
    assert Path(summary["excel_output_path"]).exists()
    assert str(summary["run_id"]).startswith("FD-RUN-")
    assert summary["status"] == "SUCCESS"
    assert summary["curve_count"] == 2
    assert summary["net_exposure_bucket_count"] > 0
    workbook = load_workbook(summary["excel_output_path"], read_only=True)
    assert "Executive Summary" in workbook.sheetnames
